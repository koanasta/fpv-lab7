import csv
import sys
from datetime import datetime

from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

CSV_PATH, OUT = sys.argv[1], sys.argv[2]

FONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
TITLE = Font(name="Arial", size=12, bold=True)
INPUT = Font(name="Arial", size=10, color="0000FF")
NOTE = Font(name="Arial", size=9, italic=True, color="666666")
HEAD_FILL = PatternFill("solid", fgColor="E7E6E6")
thin = Side(style="thin", color="999999")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def num(v):
    try:
        return float(v)
    except ValueError:
        return v


rows = list(csv.reader(open(CSV_PATH, encoding="utf-8")))
header, data = rows[0], rows[1:]
N = len(data)
LAST = N + 1

wb = Workbook()

raw = wb.active
raw.title = "CSV"
raw.append(header)
for r in data:
    out = [datetime.strptime(r[0], "%Y-%m-%d %H:%M:%S.%f")] + [num(v) for v in r[1:]]
    raw.append(out)
for c in raw[1]:
    c.font = BOLD
for row in raw.iter_rows(min_row=2, max_col=1):
    row[0].number_format = "yyyy-mm-dd hh:mm:ss"
raw.column_dimensions["A"].width = 20
raw.freeze_panes = "B2"

idx = {name: i for i, name in enumerate(header)}
cols = [("Timestamp", "Час запису"), (None, "Час від зльоту, с"),
        ("altitudeRelative", "Висота відносна, м"), ("groundSpeed", "Наземна швидкість, м/с"),
        ("airSpeed", "Повітряна швидкість, м/с"), ("roll", "Roll, град"),
        ("missionItemIndex", "Елемент місії"), ("gps.lat", "Широта"), ("gps.lon", "Довгота"),
        ("flightTime", "flightTime (QGC)"), ("distanceToHome", "Відстань від home, м"),
        (None, "Висота > 0,2 м (зліт)"), (None, "Висота > 0,05 м (до посадки)")]

d = wb.create_sheet("Дані")
d.append([c[1] for c in cols])
for i, r in enumerate(data, start=2):
    line = []
    for src, _ in cols:
        if src == "Timestamp":
            line.append(datetime.strptime(r[0], "%Y-%m-%d %H:%M:%S.%f"))
        elif src is None:
            line.append(None)
        else:
            line.append(num(r[idx[src]]))
    d.append(line)
    d.cell(i, 2, f"=ROUND((A{i}-'Показники'!$B$4)*86400,0)")
    d.cell(i, 12, f"=IF(C{i}>0.2,1,0)")
    d.cell(i, 13, f"=IF(C{i}>0.05,1,0)")
    d.cell(i, 1).number_format = "hh:mm:ss"
    for c in (8, 9):
        d.cell(i, c).number_format = "0.0000000"
for c in d[1]:
    c.font = BOLD
    c.fill = HEAD_FILL
    c.alignment = Alignment(wrap_text=True, vertical="center")
d.row_dimensions[1].height = 30
for col, w in zip("ABCDEFGHIJKLM", [11, 10, 11, 12, 12, 9, 9, 13, 13, 11, 12, 10, 12]):
    d.column_dimensions[col].width = w
d.freeze_panes = "A2"

R = lambda col: f"'Дані'!${col}$2:${col}${LAST}"

s = wb.create_sheet("Показники")
s.column_dimensions["A"].width = 38
s.column_dimensions["B"].width = 30
s.column_dimensions["C"].width = 16
s.column_dimensions["D"].width = 16
s.column_dimensions["E"].width = 46

s["A1"] = "Лабораторна 7, варіант 8 — аналіз телеметрії (CSV QGroundControl)"
s["A1"].font = TITLE

s["A3"] = "Службові"
s["A3"].font = BOLD
s["A4"] = "Момент зльоту (перший рядок з висотою > 0,2 м)"
s["B4"] = f"=INDEX({R('A')},B5)"
s["B4"].number_format = "hh:mm:ss"
s["A5"] = "Номер рядка зльоту в діапазоні"
s["B5"] = f"=MATCH(1,{R('L')},0)"
s["A6"] = "Кількість записів з висотою > 0,05 м"
s["B6"] = f"=COUNTIF({R('M')},1)"
s["A7"] = "Момент посадки (перший рядок з висотою ≤ 0,05 м)"
s["B7"] = f"=INDEX({R('A')},MATCH(1,{R('M')},0)+B6)"
s["B7"].number_format = "hh:mm:ss"

s["A9"] = "Таблиця 7.4 — Результати аналізу телеметрії"
s["A9"].font = TITLE
s["A10"], s["B10"], s["C10"], s["E10"] = "Показник", "Значення", "", "Джерело / примітка"
rows74 = [
    ("Час зльоту", "=B4", "hh:mm:ss", "CSV"),
    ("Найбільша висота, м", f"=MAX({R('C')})", "0.0", "CSV, altitudeRelative"),
    ("Найбільша швидкість (наземна), м/с", f"=MAX({R('D')})", "0.0", "CSV, groundSpeed"),
    ("Найбільша повітряна швидкість, м/с", f"=MAX({R('E')})", "0.0", "CSV, airSpeed"),
    ("Середня швидкість на горизонтальному відрізку, м/с", f"=AVERAGEIFS({R('D')},{R('G')},2,{R('D')},\">4.5\")", "0.00", "CSV, елемент місії 2, швидкість > 4,5 м/с"),
    ("Час польоту (реальний час), с", "=ROUND((B7-B4)*86400,0)", "0", "CSV пише за реальним годинником"),
    ("Час польоту (час симуляції), с", 99.2, "0.0", "З .tlog (time_boot_ms); симуляція йшла на ~0,4 від реального часу"),
    ("Початкова широта", f"='Дані'!H2", "0.0000000", "CSV, перший запис"),
    ("Початкова довгота", f"='Дані'!I2", "0.0000000", "CSV, перший запис"),
    ("Кінцева широта", f"=INDEX({R('H')},COUNTA({R('A')}))", "0.0000000", "CSV, останній запис"),
    ("Кінцева довгота", f"=INDEX({R('I')},COUNTA({R('A')}))", "0.0000000", "CSV, останній запис"),
    ("Максимальний roll за CSV, град", f"=MAX(MAX({R('F')}),-MIN({R('F')}))", "0.0", "CSV, 1 запис за секунду"),
    ("Максимальний roll за .tlog, град", 3.02, "0.0", "З .tlog (ATTITUDE, повна частота)"),
    ("Режими польоту", "STABILIZE → GUIDED → AUTO", None, "З .tlog (HEARTBEAT); у CSV режимів немає"),
    ("Стан arm до та після місії", "DISARMED → ARMED → DISARMED", None, "З .tlog (HEARTBEAT)"),
]
for k, (name, val, fmt, src) in enumerate(rows74, start=11):
    s.cell(k, 1, name)
    c = s.cell(k, 2, val)
    if fmt:
        c.number_format = fmt
    if isinstance(val, float):
        c.font = INPUT
        c.comment = Comment("Значення взято з файлу .tlog скриптом analysis/analyze.py (репозиторій fpv-lab7)", "Lab 7")
    s.cell(k, 5, src).font = NOTE
end74 = 10 + len(rows74)

top = end74 + 2
s.cell(top, 1, "Порівняння заданих і фактичних значень").font = TITLE
hdr = top + 1
for j, t in enumerate(["Параметр", "Задано", "Фактично", "Відхилення", "Примітка"], start=1):
    s.cell(hdr, j, t)
cmp_rows = [
    ("Висота, м", 40, "=B12", "=C{r}-B{r}", "Задано варіантом (H)"),
    ("Горизонтальна швидкість, м/с", 5, "=B13", "=C{r}-B{r}", "Задано варіантом (V)"),
    ("WPNAV_SPEED, см/с", "=100*B{v}", 500, "=C{r}-B{r}", "WPNAV_SPEED = 100 · V; фактичне значення з параметрів апарата"),
    ("Широта цілі", -35.3624356, "=B20", "", "Задано варіантом"),
    ("Довгота цілі", 149.1641926, "=B21", "", "Задано варіантом"),
    ("Промах до цільової точки, м", "", "", "=SQRT((RADIANS(C{la}-B{la})*6378137)^2+(RADIANS(C{lo}-B{lo})*6378137*COS(RADIANS(B{la})))^2)", "Рівнокутна апроксимація, R = 6378137 м"),
    ("Відстань старт → ціль, м", "", "", "=SQRT((RADIANS(B{la}-B18)*6378137)^2+(RADIANS(B{lo}-B19)*6378137*COS(RADIANS(B18)))^2)", "Довжина маршруту"),
]
r0 = hdr + 1
v_row = r0 + 1
la, lo = r0 + 3, r0 + 4
for k, (name, setv, act, dev, note) in enumerate(cmp_rows, start=r0):
    fill = lambda x: x.format(r=k, v=v_row, la=la, lo=lo) if isinstance(x, str) else x
    s.cell(k, 1, name)
    b = s.cell(k, 2, fill(setv))
    c = s.cell(k, 3, fill(act))
    e = s.cell(k, 4, fill(dev))
    s.cell(k, 5, note).font = NOTE
    if isinstance(setv, (int, float)):
        b.font = INPUT
    if isinstance(act, (int, float)):
        c.font = INPUT
    fmt = "0.0000000" if "цілі" in name and "Промах" not in name else "0.00"
    for x in (b, c, e):
        x.number_format = fmt
end_cmp = r0 + len(cmp_rows) - 1

for rr in list(range(10, end74 + 1)) + list(range(hdr, end_cmp + 1)):
    for cc in range(1, 6):
        cell = s.cell(rr, cc)
        cell.border = BOX
        if cell.font != INPUT and cell.font != NOTE:
            cell.font = FONT
        if rr in (10, hdr):
            cell.font = BOLD
            cell.fill = HEAD_FILL
for rr in range(3, 8):
    for cc in (1, 2):
        if s.cell(rr, cc).font != BOLD:
            s.cell(rr, cc).font = FONT
s["A3"].font = BOLD

leg = end_cmp + 2
s.cell(leg, 1, "Синій текст — значення, введені вручну (з варіанта або з .tlog); чорний — формули за даними CSV.").font = NOTE
s.cell(leg + 1, 1, "Аркуш «CSV» — вихідний файл QGroundControl без змін; аркуш «Дані» — потрібні колонки та допоміжні формули.").font = NOTE

g = wb.create_sheet("Графіки")
g["A1"] = "Графіки висоти, повітряної та наземної швидкостей від часу (дані аркуша «Дані»)"
g["A1"].font = TITLE
g["A2"] = "Етапи: зліт ≈ 0–41 с, горизонтальний політ ≈ 41–124 с, посадка ≈ 124–255 с (час від зльоту, за записами CSV)."
g["A2"].font = NOTE

x = Reference(d, min_col=2, min_row=2, max_row=LAST)
for k, (col, ttl, yt) in enumerate([(3, "Висота від часу", "Висота, м"),
                                     (5, "Повітряна швидкість від часу", "Повітряна швидкість, м/с"),
                                     (4, "Наземна швидкість від часу", "Наземна швидкість, м/с")]):
    ch = ScatterChart()
    ch.title = ttl
    ch.style = 13
    ch.x_axis.title = "Час від зльоту, с"
    ch.y_axis.title = yt
    ch.legend = None
    ch.x_axis.delete = False
    ch.x_axis.scaling.min = -10
    ch.x_axis.scaling.max = 275
    ch.x_axis.majorUnit = 25
    ch.y_axis.delete = False
    y = Reference(d, min_col=col, min_row=1, max_row=LAST)
    ser = Series(y, x, title_from_data=True)
    ser.marker.symbol = "none"
    ser.smooth = False
    ser.graphicalProperties.line.solidFill = "1F4E79"
    ser.graphicalProperties.line.width = 19050
    ch.series.append(ser)
    ch.height, ch.width = 8.5, 18
    g.add_chart(ch, f"A{4 + k * 18}")

wb.move_sheet("Показники", offset=-2)
wb.active = 0
wb.save(OUT)
print("saved", OUT, "rows", N)
